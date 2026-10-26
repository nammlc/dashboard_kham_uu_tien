"""Tạo file giả lập báo cáo Minh Lộ 04-4 (hẹn khám lại): tiêu đề lệch xuống, ô gộp 'Tuổi', bệnh nhân trải nhiều dòng,
tên entry trong zip dùng '\\' như file thật."""
import datetime as dt
import io
import zipfile

import openpyxl


def build_04_4(today: dt.date) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["A2"] = "SỞ Y TẾ", "DANH SÁCH BỆNH NHÂN HẸN KHÁM LẠI"
    hdr = ["STT", "Mã y tế", "Họ tên bệnh nhân", "Tuổi", None, "Địa chỉ", "Số thẻ BHYT", "Bác sỹ khám", "Triệu chứng",
           "Chẩn đoán", "Ngày hẹn", "Ngày lập", "Điện thoại", "Khoa hẹn", "Bác sỹ hẹn", "Đã khám"]
    for i, h in enumerate(hdr, 1):
        ws.cell(4, i, h)
    ws.merge_cells("D4:E4")
    ws.cell(5, 4, "Nam")
    ws.cell(5, 5, "Nữ")
    d = lambda n: dt.datetime.combine(today + dt.timedelta(days=n), dt.time(0))
    pts = [  # stt, mã, tên (nhiều dòng), nam, nữ, địa chỉ, ngày hẹn, ngày lập, sđt, khoa
        (1, "MYT001", ["NGUYỄN", "THỊ", "PHƯƠNG"], None, "45 tuổi", ["Xã Hoằng Quang", "Thanh Hóa"], d(5), d(0), 983123456, "Khoa Nội Tim Mạch"),
        (2, "MYT002", ["TRẦN VĂN BÌNH"], "74 tuổi", None, ["Phường Đông Sơn"], d(6), d(0), "0912 345 678", "Khoa Nội Tiêu Hóa"),
        (3, "MYT003", ["LÊ HOÀNG NAM"], "32 tháng tuổi", None, ["Xã Quảng Thắng"], d(5), d(-1), 0, "Khoa Nhi"),
        (4, "MYT004", ["PHẠM THỊ HOA"], None, "60 tuổi", ["Xã Thiệu Hóa"], d(7), d(-2), "0905111222", "Khoa Nội Thần Kinh"),
        (5, "MYT005", ["ĐỖ THỊ NGA"], None, "51 tuổi", ["Xã Đông Hưng"], d(15), d(0), "0372984602", "Khoa Sản"),
    ]
    r = 6
    for stt, ma, tens, nam, nu, dcs, hen, lap, sdt, khoa in pts:
        for k in range(max(len(tens), len(dcs))):
            if k == 0:
                ws.cell(r, 1, stt); ws.cell(r, 2, ma)
                if nam: ws.cell(r, 4, nam)
                if nu: ws.cell(r, 5, nu)
                ws.cell(r, 11, hen); ws.cell(r, 12, lap); ws.cell(r, 13, sdt); ws.cell(r, 14, khoa)
            if k < len(tens): ws.cell(r, 3, tens[k])
            if k < len(dcs): ws.cell(r, 6, dcs[k])
            r += 1
    for row in ws.iter_rows(min_row=6):
        for c in row:
            if isinstance(c.value, dt.datetime):
                c.number_format = "dd/mm/yyyy"
    buf = io.BytesIO()
    wb.save(buf)
    zin = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in zin.namelist():
            zout.writestr(n.replace("/", "\\"), zin.read(n))
    return out.getvalue()
